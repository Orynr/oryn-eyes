import sys, os
sys.path.insert(0, os.path.dirname(__file__))

import openpyxl
import excel_export
import sheets_export

echecs = []

def check(label, condition):
    status = "OK" if condition else "FAIL"
    if not condition:
        echecs.append(label)
    print(f"[{status}] {label}")


donnees = {
    "cpu_amd": sheets_export.enrichir_listings([
        {"modele": "Ryzen 5 3600", "prix": 25, "plateforme": "Vinted",
         "lot_type": "normal", "score_vendeur": 4.8, "etat": "Bon",
         "url": "http://x/1", "horodatage": "2026-08-20T10:00:00"},
        {"modele": "Ryzen 5 3600", "prix": 45, "plateforme": "eBay",
         "lot_type": "normal", "score_vendeur": 4.2, "etat": "TBE",
         "url": "http://x/2", "horodatage": "2026-08-20T11:00:00"},
    ]),
    "gpu": sheets_export.enrichir_listings([
        {"modele": "RTX 3070", "prix": 175, "plateforme": "Vinted",
         "lot_type": "normal", "score_vendeur": 4.9, "etat": "Bon",
         "url": "http://x/3", "horodatage": "2026-08-20T12:00:00"},
    ]),
    "ram": [],  # catégorie sans annonce -> ne doit pas planter
}

os.chdir("/tmp")
chemin = excel_export.generer_backup(donnees, chemin="test_backup.xlsx")

check("le fichier existe", os.path.exists(chemin))

wb = openpyxl.load_workbook(chemin)
check("onglet Occasion présent", "Occasion" in wb.sheetnames)
check("onglet Resume présent", "Resume" in wb.sheetnames)

ws_occ = wb["Occasion"]
check("en-têtes corrects", ws_occ.cell(row=1, column=1).value == "categorie")
check("3 lignes de données (2 CPU + 1 GPU)", ws_occ.max_row == 4)  # 1 entête + 3 lignes

ws_res = wb["Resume"]
check("3 catégories dans le résumé (dont RAM vide)", ws_res.max_row == 4)  # entête + 3 cat
ram_row = [ws_res.cell(row=r, column=1).value for r in range(2, 5)]
check("RAM apparaît avec des tirets, pas de crash", "RAM" in ram_row)

os.remove(chemin)

print("\n" + "=" * 50)
if echecs:
    print(f"ECHECS ({len(echecs)}):")
    for e in echecs:
        print(f"  - {e}")
    sys.exit(1)
else:
    print("TOUT PASSE.")
